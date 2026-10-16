"""Bounded, deterministic conversation evidence. Never run Excel or evaluate formulas."""
import csv
import datetime
import hashlib
import io
import math
import re
import zipfile
from pathlib import Path

VERSION = 'chat-spreadsheet-1'
MAX_BYTES = 12 * 1024 * 1024
MAX_CELLS = 6000


def column_name(index):
    name = ''
    while index:
        index, remainder = divmod(index - 1, 26)
        name = chr(65 + remainder) + name
    return name


def cell_position(value):
    match = re.fullmatch(r'([A-Z]{1,3})([1-9]\d{0,6})', str(value).upper())
    if not match:
        raise ValueError('请输入单元格范围，例如 A1:F50')
    column = 0
    for char in match[1]:
        column = column * 26 + ord(char) - 64
    row = int(match[2])
    if row > 1048576 or column > 16384:
        raise ValueError('单元格范围超过 Excel 工作表边界')
    return row, column


def selected_ranges(values, rows, columns):
    if not values:
        return [(1, 1, max(1, min(rows, 100)), max(1, min(columns, 30)))]
    if not isinstance(values, list) or len(values) > 3:
        raise ValueError('一次最多选择 3 处单元格范围')
    result = []
    for value in values:
        if not isinstance(value, str) or value.count(':') > 1:
            raise ValueError('请输入单元格范围，例如 A1:F50')
        parts = value.split(':')
        start, end = cell_position(parts[0]), cell_position(parts[-1])
        if end[0] < start[0] or end[1] < start[1]:
            raise ValueError('单元格范围的结束位置应在起始位置之后')
        if end[0] - start[0] + 1 > 200 or end[1] - start[1] + 1 > 60:
            raise ValueError('每处最多读取 200 行、60 列，请缩小选区')
        result.append((start[0], start[1], end[0], end[1]))
    if sum((r2-r1+1)*(c2-c1+1) for r1,c1,r2,c2 in result) > MAX_CELLS:
        raise ValueError('本次范围超过 6000 个单元格，请缩小选区')
    return result


def plain_value(value):
    if value is None:
        return ''
    if isinstance(value, (datetime.date, datetime.datetime, datetime.time)):
        return value.isoformat()
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, float):
        if not math.isfinite(value):
            return str(value)
        return str(int(value)) if value.is_integer() else str(value)
    return str(value)


def read_spreadsheet(filename, sheet=None, ranges=None):
    path = Path(filename)
    if not path.is_file() or not 0 < path.stat().st_size <= MAX_BYTES:
        raise ValueError('表格为空或超过 12 MB，请选择较小的文件')
    binary = path.read_bytes()
    fingerprint = hashlib.sha256(binary).hexdigest()
    extension = path.suffix.lower()
    warnings, workbook, closer = [], None, lambda: None
    if extension == '.csv':
        try:
            text = binary.decode('utf-8-sig')
        except UnicodeDecodeError:
            try:
                text = binary.decode('gb18030')
                warnings.append('CSV 按 GB18030 解码；请核对中文与单位。')
            except UnicodeDecodeError:
                raise ValueError('CSV 编码无法识别，请另存为 UTF-8 CSV') from None
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        csv.field_size_limit(1024 * 1024)
        parsed, columns = [], 0
        try:
            for row in csv.reader(io.StringIO(text, newline=''), dialect=dialect, strict=True):
                if len(parsed) >= 100000 or len(row) > 500:
                    raise ValueError('CSV 超过 10 万行或 500 列，请另存所需工作表范围')
                parsed.append(row)
                columns = max(columns, len(row))
        except csv.Error:
            raise ValueError('CSV 的引号或分隔符无效，请重新导出 CSV') from None
        # CSV has no embedded sheet name. Use a stable label instead of Zotero's
        # internal attachment-storage filename, which can change on import.
        sheets = [dict(name='CSV', rows=len(parsed), columns=columns, hidden=False)]
        def get_rows(_name, r1, c1, r2, c2):
            for r in range(r1, r2 + 1):
                row = parsed[r-1] if r <= len(parsed) else []
                yield r, [row[c-1] if c <= len(row) else '' for c in range(c1, c2+1)]
    elif extension == '.xlsx':
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ValueError('Excel 解析工具缺失，请修复完整安装包或先另存为 CSV') from None
        try:
            with zipfile.ZipFile(io.BytesIO(binary)) as archive:
                entries = archive.infolist()
                if len(entries) > 2000 or sum(entry.file_size for entry in entries) > 80 * 1024 * 1024:
                    raise ValueError('Excel 解压内容过大，请另存所需工作表')
                if any(entry.flag_bits & 1 for entry in entries):
                    raise ValueError('不支持加密的 Excel，请另存为未加密 XLSX')
            workbook = load_workbook(io.BytesIO(binary), read_only=True, data_only=False, keep_links=False)
            closer = workbook.close
            if len(workbook.sheetnames) > 100:
                closer()
                raise ValueError('工作表超过 100 张，请另存所需工作表')
            sheets = [dict(name=tab.title, rows=tab.max_row or 0, columns=tab.max_column or 0,
                           hidden=tab.sheet_state != 'visible') for tab in workbook.worksheets]
            def get_rows(name, r1, c1, r2, c2):
                for number, row in enumerate(workbook[name].iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2, values_only=True), r1):
                    yield number, list(row)
        except (zipfile.BadZipFile, KeyError, OSError):
            raise ValueError('文件不是可读取的 XLSX，请重新另存为 XLSX 或 CSV') from None
        warnings.append('Excel 公式按原式显示，未计算；不读取外部链接，也不执行宏。')
    elif extension == '.xls':
        try:
            import xlrd
        except ImportError:
            raise ValueError('旧版 XLS 解析工具缺失，请修复完整安装包或另存为 XLSX/CSV') from None
        try:
            workbook = xlrd.open_workbook(file_contents=binary, on_demand=True)
            closer = workbook.release_resources
            if workbook.nsheets > 100:
                closer()
                raise ValueError('工作表超过 100 张，请另存所需工作表')
            sheets = [dict(name=tab.name, rows=tab.nrows, columns=tab.ncols, hidden=tab.visibility != 0) for tab in workbook.sheets()]
            def get_rows(name, r1, c1, r2, c2):
                tab = workbook.sheet_by_name(name)
                for r in range(r1, r2 + 1):
                    values = []
                    for c in range(c1, c2+1):
                        cell = tab.cell(r-1, c-1) if r <= tab.nrows and c <= tab.ncols else None
                        value = '' if cell is None else cell.value
                        if cell and cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(value, workbook.datemode)
                        elif cell and cell.ctype == xlrd.XL_CELL_BOOLEAN:
                            value = bool(value)
                        elif cell and cell.ctype == xlrd.XL_CELL_ERROR:
                            value = xlrd.error_text_from_code.get(value, '#ERROR')
                        values.append(value)
                    yield r, values
        except xlrd.XLRDError:
            raise ValueError('文件不是可读取的旧版 XLS，请另存为 XLSX 或 CSV') from None
        warnings.append('XLS 公式仅能读取文件中上次保存的值，未重新计算；不执行宏。')
    else:
        raise ValueError('请选择 XLSX、XLS 或 CSV 表格')
    try:
        if not sheets:
            raise ValueError('文件中没有可读取的工作表')
        selected = next((value for value in sheets if value['name'] == sheet), None) if sheet else next((value for value in sheets if not value['hidden']), sheets[0])
        if not selected:
            raise ValueError('指定的工作表已不存在，请重新选择工作表')
        chosen = selected_ranges(ranges, selected['rows'], selected['columns'])
        if selected['hidden']:
            warnings.append('当前读取的是隐藏工作表。')
        sections, character_count, shortened = [], 0, False
        for r1, c1, r2, c2 in chosen:
            cells, formulas = [], []
            for r, values in get_rows(selected['name'], r1, c1, r2, c2):
                row = []
                for c, value in enumerate(values, c1):
                    raw = plain_value(value)
                    if raw.startswith('=') and extension == '.xlsx':
                        formulas.append(column_name(c) + str(r))
                    if len(raw) > 1000:
                        shortened = True
                        raw = raw[:1000] + '…（单元格内容已截取）'
                    character_count += len(raw)
                    if character_count > 250000:
                        raise ValueError('所选单元格文本超过 25 万字，请缩小选区')
                    row.append(raw)
                cells.append(dict(row=r, values=row))
            address = f'{column_name(c1)}{r1}:{column_name(c2)}{r2}'
            sections.append(dict(range=address, startColumn=c1, endColumn=c2, cells=cells, formulaCells=formulas))
        truncated = shortened or not ranges and (selected['rows'] > 100 or selected['columns'] > 30)
        if truncated:
            warnings.append('仅显示所选范围的内容；可选择工作表与单元格范围继续读取，不能据此声称已读完整文件。')
        return dict(version=VERSION, hash=fingerprint, sheets=sheets, selectedSheet=selected['name'],
                    range=sections[0]['range'], cells=sections[0]['cells'], sections=sections,
                    truncated=bool(truncated), warnings=warnings)
    finally:
        closer()

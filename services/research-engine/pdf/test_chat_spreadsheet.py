"""Real workbook/CSV bytes; no Excel process, network or model."""
import csv
import tempfile
import unittest
import zipfile
from pathlib import Path
from openpyxl import Workbook
from chat_spreadsheet import read_spreadsheet, selected_ranges
from assets import main


class ConversationSpreadsheetTests(unittest.TestCase):
    def test_workbook_sheets_ranges_formula_and_dates_use_actual_cells(self):
        import datetime
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / '测量.xlsx'
            book = Workbook()
            book.active.title = '概览'
            book.active.append(['说明', '只是索引，不是实验'])
            data = book.create_sheet('田间实验')
            data.append(['处理', '均值', '重复数', '公式', '采样日期'])
            data.append(['对照', 7.5, 4, '=B2*C2', datetime.date(2026, 10, 15)])
            hidden = book.create_sheet('隐藏数据')
            hidden.sheet_state = 'hidden'
            hidden['A1'] = '不能默认读这个表'
            book.save(path)
            value = read_spreadsheet(path, '田间实验', ['B1:E2'])
            self.assertEqual([sheet['name'] for sheet in value['sheets']], ['概览', '田间实验', '隐藏数据'])
            self.assertTrue(value['sheets'][2]['hidden'])
            self.assertEqual(value['selectedSheet'], '田间实验')
            self.assertEqual(value['range'], 'B1:E2')
            self.assertEqual(value['cells'][1], dict(row=2, values=['7.5', '4', '=B2*C2', '2026-10-15T00:00:00']))
            self.assertEqual(value['sections'][0]['formulaCells'], ['D2'])
            self.assertEqual(value['sections'][0]['startColumn'], 2)
            self.assertEqual(len(value['hash']), 64)
            self.assertIn('未计算', ''.join(value['warnings']))
            self.assertEqual(read_spreadsheet(path)['selectedSheet'], '概览')
            with self.assertRaisesRegex(ValueError, '工作表已不存在'):
                read_spreadsheet(path, '已经删除')

    def test_csv_handles_quoted_multiline_chinese_encoding_and_default_coverage(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / '数据.csv'
            with path.open('w', encoding='utf-8-sig', newline='') as file:
                writer = csv.writer(file)
                writer.writerow(['处理', '含义', '观测值'])
                writer.writerow(['A', '第一行\n第二行,保留逗号', '0.850'])
                for number in range(2, 121):
                    writer.writerow([f'处理{number}', '', str(number)])
            value = read_spreadsheet(path)
            self.assertEqual(value['cells'][1]['values'], ['A', '第一行\n第二行,保留逗号', '0.850'])
            self.assertEqual(value['sheets'][0]['rows'], 121)
            self.assertTrue(value['truncated'])
            self.assertEqual(len(value['cells']), 100)
            selected = read_spreadsheet(path, ranges=['C110:C112'])
            self.assertEqual(selected['cells'][0], dict(row=110, values=['109']))
            previous = value['hash']
            path.write_text('处理,结果\n对照,7.8\n', encoding='gb18030')
            changed = read_spreadsheet(path)
            self.assertNotEqual(previous, changed['hash'])
            self.assertEqual(changed['cells'][1]['values'], ['对照', '7.8'])
            self.assertIn('GB18030', ''.join(changed['warnings']))

    def test_real_legacy_xls_fixture_and_asset_engine_route(self):
        fixture = Path(__file__).resolve().parent.parent / 'test' / 'fixtures' / 'chat-lab-data.xls'
        with tempfile.TemporaryDirectory() as directory:
            value = main(dict(operation='assets-chat-spreadsheet', file=str(fixture), cacheRoot=directory,
                              sheet='旧版实验', ranges=['A1:B2']))
            self.assertEqual(value['cells'][1]['values'], ['对照', '7.5'])
            self.assertIn('上次保存的值', ''.join(value['warnings']))

    def test_invalid_or_excessive_ranges_and_fake_workbooks_are_rejected(self):
        for ranges in [['A0:B1'], ['A5:B2'], ['XFE1:XFE2'], ['A1:Z300'], ['A1:BH101'], ['A1'] * 4]:
            with self.assertRaises(ValueError):
                selected_ranges(ranges, 1000, 100)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / '错误.xlsx'
            path.write_text('<html>not a workbook</html>')
            with self.assertRaisesRegex(ValueError, '不是可读取的 XLSX'):
                read_spreadsheet(path)
            with zipfile.ZipFile(path, 'w') as file:
                for number in range(2001):
                    file.writestr(str(number), '')
            with self.assertRaisesRegex(ValueError, '解压内容过大'):
                read_spreadsheet(path)


if __name__ == '__main__':
    unittest.main()

import tempfile, unittest, zipfile
from pathlib import Path
from unittest.mock import patch
from document import export, dataset, qn, JOURNAL_LAYOUTS
from read_export import export_read
from lxml import etree
from docx import Document
from docx.shared import Mm, Pt
from copy import deepcopy

class DocumentTest(unittest.TestCase):
 def test_execution_plan_task_keeps_estimated_hours_with_preceding_task_fields(self):
  with tempfile.TemporaryDirectory() as tmp:
   fields=['执行：核对测量条件。','交付：复核记录。','检查：来源与页码可追溯。','预计：24–55 小时。']
   project=dict(title='执行计划',sections=[dict(id='overview',title='评估概况'),dict(id='task',title='测量核对')],blocks=[dict(id='overview-'+str(i),type='paragraph',sectionID='overview',text='概况段落 '+str(i)) for i in range(4)]+[dict(id='p'+str(i),type='paragraph',sectionID='task',text=text) for i,text in enumerate(fields)])
   before=deepcopy(project);result=export(dict(directory=tmp,operation='manuscript-plan-export',project=project))
   paragraphs=Document(result['path']).paragraphs
   task=next(i for i,p in enumerate(paragraphs) if p.text=='测量核对')
   self.assertEqual([p.text for p in paragraphs[task+1:]],fields)
   self.assertEqual([p.paragraph_format.keep_with_next for p in paragraphs[task:]],[True,True,True,True,False])
   self.assertTrue(all(p.paragraph_format.keep_together for p in paragraphs[task:]))
   self.assertTrue(all(p.paragraph_format.keep_with_next is not True for p in paragraphs if p.text.startswith('概况段落')))
   self.assertEqual(project,before)

 def test_table_geometry_and_reference_styles(self):
  with tempfile.TemporaryDirectory() as tmp:
   file=Path(tmp)/'data.csv';file.write_text('方法,准确率\n甲,95%\n乙,96%',encoding='utf-8');data=dataset(file)
   self.assertEqual(data['rows'][2][1],'96%')
   rows=[[dict(text='实验比较',rowspan=1,colspan=2),dict(text='',hidden=True)], [dict(text='甲',rowspan=1,colspan=1),dict(text='95%',rowspan=1,colspan=1,align='center')]]
   b=dict(id='table-one',type='table',sectionID='s',table=dict(rows=rows,width=158,widths=[70,88],heights=[10,12],style='grid',caption='比较',border=1))
   result=export(dict(directory=tmp,project=dict(title='验收',sections=[dict(id='s',title='结果')],blocks=[b,dict(id='ref',type='crossref',sectionID='s',targetID=b['id'],text='表 1')]),bibliography=[]))
   with zipfile.ZipFile(result['path']) as z:
    root=etree.fromstring(z.read('word/document.xml'))
    self.assertEqual(root.find('.//'+qn('w:gridSpan')).get(qn('w:val')),'2')
    self.assertEqual(root.find('.//'+qn('w:tcBorders')+'/'+qn('w:top')).get(qn('w:sz')),'8')
    self.assertIsNotNone(root.find('.//'+qn('w:trHeight')))
    self.assertIn('REF table_table_one',z.read('word/document.xml').decode())
    self.assertIn('宋体',z.read('word/styles.xml').decode())
    self.assertIn('Cambria Math',z.read('word/settings.xml').decode())
    self.assertIn('word/footer1.xml',z.namelist())

 def test_journal_presets_change_real_word_geometry_and_styles_without_rewriting(self):
  with tempfile.TemporaryDirectory() as tmp:
   for format_id, layout in JOURNAL_LAYOUTS.items():
    project=dict(title='期刊稿件',journalFormat=dict(id=format_id),sections=[dict(id='s',title='方法')],blocks=[dict(id='p',type='paragraph',sectionID='s',text='保留原文：测量值为 2.8%，未完成独立复测。')]);before=deepcopy(project)
    result=export(dict(directory=tmp,project=project))
    with zipfile.ZipFile(result['path']) as archive:
     xml=etree.fromstring(archive.read('word/document.xml'));styles=etree.fromstring(archive.read('word/styles.xml'));normal=next(node for node in styles.findall(qn('w:style')) if node.get(qn('w:styleId'))=='Normal')
     self.assertEqual(normal.find(qn('w:rPr')+'/'+qn('w:rFonts')).get(qn('w:ascii')),layout['font'])
     self.assertEqual(normal.find(qn('w:rPr')+'/'+qn('w:sz')).get(qn('w:val')),str(layout['fontSize']*2))
     self.assertEqual(normal.find(qn('w:pPr')+'/'+qn('w:spacing')).get(qn('w:line')),str(round(layout['lineSpacing']*240)))
     sections=xml.findall('.//'+qn('w:sectPr'));self.assertEqual(sections[-1].find(qn('w:cols')).get(qn('w:num')),str(layout['columns']))
     self.assertAlmostEqual(int(sections[-1].find(qn('w:pgMar')).get(qn('w:left'))),layout['margin']*56.6929,delta=1)
     self.assertIn(project['blocks'][0]['text'],''.join(xml.itertext()))
     if layout['columns']==2:self.assertEqual(sections[0].find(qn('w:cols')).get(qn('w:num')),'1')
    self.assertEqual(project,before)

 def test_custom_word_reference_remains_authoritative_when_journal_citation_format_changes(self):
  with tempfile.TemporaryDirectory() as tmp:
   reference=Path(tmp)/'custom.docx';template=Document();template.styles['Normal'].font.name='Courier New';template.styles['Normal'].font.size=Pt(14);template.styles['Normal'].paragraph_format.line_spacing=1.15;template.sections[0].left_margin=Mm(32);template.save(reference)
   result=export(dict(directory=tmp,referenceDoc=str(reference),project=dict(title='自定义',journalFormat=dict(id='ieee'),sections=[dict(id='s',title='方法')],blocks=[dict(id='p',type='paragraph',sectionID='s',text='保留自定义模板与正文。')])))
   doc=Document(result['path']);self.assertEqual(doc.styles['Normal'].font.name,'Courier New');self.assertEqual(doc.styles['Normal'].font.size.pt,14);self.assertAlmostEqual(doc.sections[0].left_margin.mm,32,delta=.02);self.assertEqual(len(doc.sections),1)

 def test_two_column_tables_fit_the_selected_column_without_mutating_source_geometry(self):
  with tempfile.TemporaryDirectory() as tmp:
   block=dict(id='t',type='table',sectionID='s',table=dict(rows=[[dict(text='原始表格'),dict(text='93%')]],width=158,widths=[79,79],heights=[10],style='grid'))
   project=dict(title='双栏稿件',journalFormat=dict(id='ieee'),sections=[dict(id='s',title='结果')],blocks=[block]);before=deepcopy(project);result=export(dict(directory=tmp,project=project))
   with zipfile.ZipFile(result['path']) as archive:
    xml=etree.fromstring(archive.read('word/document.xml'));width=int(xml.find('.//'+qn('w:tblW')).get(qn('w:w')));self.assertLessEqual(width,round(82.25*56.693));self.assertIn('原始表格',''.join(xml.itertext()))
   self.assertEqual(project,before)

 def test_read_export_writes_a_new_file_with_its_header_and_keeps_four_block_tasks_together(self):
  with tempfile.TemporaryDirectory() as tmp:
   task=['完成独立对照\n1. 准备数据\n2. 运行对照','交付物 对照记录','验收条件 两组结果可追溯','预计 2–4 小时（按现有数据）']
   project=dict(title='速读',sections=[dict(id='c',title='方法'),dict(id='t',title='任务 1 · 独立对照')],blocks=[dict(id='c'+str(i),type='paragraph',sectionID='c',text='证据要点 · 有依据\n要点 '+str(i)) for i in range(3)]+[dict(id='t'+str(i),type='paragraph',sectionID='t',text=text) for i,text in enumerate(task)])
   before=deepcopy(project);first=export_read(dict(directory=tmp,operation='manuscript-read-export',project=project));second=export_read(dict(directory=tmp,operation='manuscript-read-export',project=project))
   self.assertEqual(Path(first['path']).name,'assembly-read.docx');self.assertEqual(first['kind'],'assembly-read');self.assertNotEqual(first['path'],second['path']);self.assertTrue(Path(first['path']).is_file() and Path(second['path']).is_file());self.assertEqual(project,before)
   doc=Document(first['path']);self.assertIn('速读与任务计划',[p.text for s in doc.sections for p in s.header.paragraphs])
   paragraphs=doc.paragraphs;start=next(i for i,p in enumerate(paragraphs) if p.text=='任务 1 · 独立对照');chapter=next(i for i,p in enumerate(paragraphs) if p.text=='方法')
   self.assertEqual([p.paragraph_format.keep_with_next for p in paragraphs[start:]],[True,True,True,True,False])
   self.assertTrue(all(p.paragraph_format.keep_with_next is not True for p in paragraphs[chapter:start]),'chapter sections are not kept together')
   plan=Document(export(dict(directory=tmp,operation='manuscript-plan-export',project=project))['path']);self.assertIn('执行方案',[p.text for s in plan.sections for p in s.header.paragraphs])

 def test_read_export_never_replaces_an_existing_file(self):
  with tempfile.TemporaryDirectory() as tmp:
   folder=Path(tmp)/'fixed';folder.mkdir();(folder/'assembly-read.docx').write_text('old',encoding='utf-8')
   project=dict(title='速读',sections=[dict(id='c',title='方法')],blocks=[dict(id='c',type='paragraph',sectionID='c',text='要点')])
   with patch('read_export.tempfile.mkdtemp',return_value=str(folder)):
    with self.assertRaises(ValueError):export_read(dict(directory=tmp,operation='manuscript-read-export',project=project))
   self.assertEqual((folder/'assembly-read.docx').read_text(encoding='utf-8'),'old')

if __name__=='__main__':unittest.main()

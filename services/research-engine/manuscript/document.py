"""Deterministic manuscript IO. No network or model execution."""
import csv, hashlib, io, json, re, subprocess, sys, tempfile, zipfile
from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from lxml import etree

def element(tag, **attrs):
    el = OxmlElement(tag)
    for key, value in attrs.items(): el.set(qn('w:'+key), str(value))
    return el

def field(p, instruction, text):
    r = p.add_run()._r
    r.append(element('w:fldChar', fldCharType='begin'))
    i = element('w:instrText'); i.text = instruction; r.append(i)
    r.append(element('w:fldChar', fldCharType='separate'))
    t = element('w:t'); t.text = str(text); r.append(t)
    r.append(element('w:fldChar', fldCharType='end'))

def paragraph_alignment(paragraph, block):
    alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
                 'right': WD_ALIGN_PARAGRAPH.RIGHT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}.get(block.get('textAlign'))
    if alignment is not None:
        paragraph.alignment = alignment
    if block.get('lineSpacing') in (1, 1.15, 1.5, 1.8, 2):
        paragraph.paragraph_format.line_spacing = block['lineSpacing']
    indent = block.get('indent')
    if isinstance(indent, (int, float)) and 0 <= indent <= 30:
        paragraph.paragraph_format.left_indent = Mm(indent)

def color_hex(value):
    color = str(value or '').lstrip('#')
    if color.startswith('rgb'):
        values = re.findall(r'\d+', color)[:3]
        color = ''.join(format(min(255, int(v)), '02X') for v in values)
    return color if re.fullmatch(r'[0-9a-fA-F]{6}', color) else None

INLINE_MATH = re.compile(r'\$\$([\s\S]+?)\$\$|\\\[([\s\S]+?)\\\]|\\\(([\s\S]+?)\\\)|(?<!\\)\$(?!\$)([^\n$]+?)(?<!\\)\$(?!\$)')


def write_text_runs(paragraph, specs):
    for spec in specs:
        run = paragraph.add_run(spec.get('text', ''))
        for name in ('bold', 'italic', 'underline'):
            if name in spec: setattr(run, name, bool(spec[name]))
        if 'strike' in spec: run.font.strike = bool(spec['strike'])
        if spec.get('sup'): run.font.superscript = True
        elif spec.get('sub'): run.font.subscript = True
        size = spec.get('fontSize')
        if isinstance(size, (int, float)) and 6 <= size <= 72: run.font.size = Pt(size)
        color = color_hex(spec.get('color'))
        if color: run.font.color.rgb = __import__('docx.shared', fromlist=['RGBColor']).RGBColor.from_string(color)
        highlight = color_hex(spec.get('highlight'))
        if highlight: run._element.get_or_add_rPr().append(element('w:shd', fill=highlight, val='clear'))
        if spec.get('fontFamily'):
            font = spec['fontFamily'].strip('"\'')
            run.font.name = font
            run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), font)


def text_run_range(specs, start, end):
    offset = 0
    for spec in specs:
        value = spec.get('text', '')
        left, right = max(start - offset, 0), min(end - offset, len(value))
        if right > left:
            yield dict(spec, text=value[left:right])
        offset += len(value)


def current_text_runs(specs, text):
    """Retain trusted formatting while the actual saved text owns the content."""
    parts = specs or [dict(text=text)]
    source = ''.join(spec.get('text', '') for spec in parts)
    if source == text:
        return parts
    # Legacy cell edits sometimes left their previous runs in storage. Keep
    # exact unchanged prefix/suffix formatting, never any superseded characters.
    prefix = 0
    while prefix < min(len(source), len(text)) and source[prefix] == text[prefix]:
        prefix += 1
    suffix = 0
    while suffix < min(len(source)-prefix, len(text)-prefix) and source[-1-suffix] == text[-1-suffix]:
        suffix += 1
    result = list(text_run_range(parts, 0, prefix))
    middle = text[prefix:len(text)-suffix if suffix else len(text)]
    if middle:
        result.append(dict(text=middle))
    if suffix:
        result.extend(text_run_range(parts, len(source)-suffix, len(source)))
    return result


def write_runs(paragraph, specs, fallback=None, math_context=None):
    parts = specs or [dict(text=fallback or '')]
    if fallback is not None and ''.join(spec.get('text', '') for spec in parts) != fallback:
        parts = current_text_runs(parts, fallback)
        if math_context is not None:
            math_context.setdefault('textWarnings', []).append(dict(reason='saved-run-text-mismatch', usedCurrentText=True))
    source = ''.join(spec.get('text', '') for spec in parts)
    matches = list(INLINE_MATH.finditer(source)) if math_context else []
    if not matches:
        write_text_runs(paragraph, parts)
        return
    offset = 0
    for match in matches:
        write_text_runs(paragraph, text_run_range(parts, offset, match.start()))
        latex = next(value for value in match.groups() if value is not None)
        try:
            nodes = formula(latex, math_context.get('pandoc'), math_context['folder'])
            for node in nodes:
                paragraph._p.append(node)
        except Exception as error:
            # A failed conversion remains editable source text, including its
            # formatting and delimiters. The export reports this explicitly.
            write_text_runs(paragraph, text_run_range(parts, match.start(), match.end()))
            math_context['warnings'].append(dict(source=match.group(0)[:400], error=str(error)[:400]))
        offset = match.end()
    write_text_runs(paragraph, text_run_range(parts, offset, len(source)))

def list_runs(text, specs):
    start = 0
    for line in text.split('\n'):
        end, offset, parts = start + len(line), 0, []
        for spec in specs:
            value = spec.get('text', '')
            left, right = max(start-offset, 0), min(end-offset, len(value))
            if right > left: parts.append(dict(spec, text=value[left:right]))
            offset += len(value)
        yield line, parts
        start = end + 1

def default_reference(path):
    doc = Document()
    for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption']:
        s = doc.styles[name]; s.font.name = 'Times New Roman'; s.font.size = Pt(11 if name == 'Normal' else 14)
        s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体' if name == 'Normal' else '黑体')
    sect = doc.sections[0]; sect.page_width=Mm(210); sect.page_height=Mm(297)
    sect.top_margin=sect.bottom_margin=Mm(25.4); sect.left_margin=sect.right_margin=Mm(26)
    sect.header.paragraphs[0].text='论文装配 · 研究稿'
    p=sect.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; field(p,' PAGE ','1')
    math = OxmlElement('m:mathPr'); font = OxmlElement('m:mathFont'); font.set(qn('m:val'),'Cambria Math'); math.append(font); doc.settings.element.append(math)
    doc.save(path)

EXPORT_HEADERS = {'manuscript-plan-export': '执行方案', 'manuscript-read-export': '速读与任务计划'}
KEEP_TASK_OPERATIONS = ('manuscript-plan-export', 'manuscript-read-export')


def academic_default_styles(doc, header_text='研究稿'):
    # Only the bundled default is normalized. A user-selected reference.docx
    # remains authoritative for fonts, colors, borders, headers and footers.
    for name in ['Normal', 'Title', 'Caption'] + ['Heading '+str(i) for i in range(1, 10)]:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.color.rgb = RGBColor(0, 0, 0)
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.get_or_add_rFonts()
        fonts.set(qn('w:eastAsia'), '宋体' if name in ('Normal', 'Caption') else '黑体')
        color = rpr.find(qn('w:color'))
        for attr in ('themeColor', 'themeTint', 'themeShade'):
            color.attrib.pop(qn('w:'+attr), None)
        if name == 'Title' or name.startswith('Heading '):
            ppr = style.element.find(qn('w:pPr'))
            if ppr is not None:
                for border in ppr.findall(qn('w:pBdr')):
                    ppr.remove(border)
    for section in doc.sections:
        for header in (section.header, section.first_page_header, section.even_page_header):
            if header.is_linked_to_previous:
                continue
            for paragraph in header.paragraphs:
                paragraph.text = header_text
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor.from_string('777777')
                    run.font.name = 'Times New Roman'
                    run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')


JOURNAL_LAYOUTS = {
    'research': dict(font='Times New Roman', fontSize=11, lineSpacing=1.5, columns=1, margin=25.4),
    'nature': dict(font='Arial', fontSize=11, lineSpacing=2, columns=1, margin=25.4),
    'physical-review': dict(font='Times New Roman', fontSize=12, lineSpacing=2, columns=1, margin=25.4),
    'ieee': dict(font='Times New Roman', fontSize=10, lineSpacing=1, columns=2, margin=19),
}


def section_columns(section, count):
    cols = section._sectPr.find(qn('w:cols'))
    if cols is None:
        cols = element('w:cols')
        section._sectPr.append(cols)
    for child in list(cols):
        cols.remove(child)
    cols.set(qn('w:num'), str(count))
    cols.set(qn('w:equalWidth'), '1')
    cols.set(qn('w:space'), '425')  # 7.5 mm between columns


def apply_journal_layout(doc, value):
    """Apply only known presets. A custom reference.docx remains authoritative."""
    if not value:
        return None
    layout = JOURNAL_LAYOUTS.get(value.get('id'))
    if not layout:
        raise ValueError('未知期刊稿件格式，请重新选择')
    for name in ['Normal', 'Title', 'Caption', 'List Bullet', 'List Number'] + ['Heading '+str(i) for i in range(1, 10)]:
        style = doc.styles[name]
        style.font.name = layout['font']
        style.font.size = Pt(layout['fontSize'] if name in ('Normal', 'Caption', 'List Bullet', 'List Number') else 18 if name == 'Title' else 13 if name == 'Heading 1' else 12)
        style.paragraph_format.line_spacing = layout['lineSpacing']
        if name in ('Normal', 'Caption', 'List Bullet', 'List Number'):
            style.paragraph_format.space_after = Pt(0)
    for section in doc.sections:
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(layout['margin'])
        section_columns(section, layout['columns'])
    return layout

def dataset(file):
    p=Path(file)
    if p.suffix.lower()=='.csv':
        try: text=p.read_text(encoding='utf-8-sig')
        except UnicodeDecodeError: text=p.read_text(encoding='gb18030')
        rows=list(csv.reader(io.StringIO(text)))
    elif p.suffix.lower()=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(p,read_only=True,data_only=True)
        rows=[list(r) for r in book.active.iter_rows(values_only=True)]; book.close()
    else: raise ValueError('请选择 CSV 或 XLSX 数据文件')
    rows=[r for r in rows if any(x is not None and str(x)!='' for x in r)]
    if not rows or len(rows)>100 or max(map(len,rows))>30: raise ValueError('数据需包含 1–100 行、1–30 列；请先选择实验所需范围')
    cols=max(map(len,rows)); rows=[[str(v) if v is not None else '' for v in r]+['']*(cols-len(r)) for r in rows]
    return dict(rows=rows,file=str(p.resolve()),hash=hashlib.sha256(p.read_bytes()).hexdigest(),sheet=p.stem)

def formula(latex,pandoc,folder):
    if not pandoc or not Path(pandoc).is_file(): raise ValueError('公式转换需要配置 Pandoc')
    out=Path(folder)/'formula.docx'
    run=subprocess.run([pandoc,'-f','markdown+tex_math_dollars','-t','docx','-o',str(out)],input='$$\n'+latex+'\n$$',text=True,capture_output=True,encoding='utf-8',timeout=30)
    if run.returncode or 'Could not convert' in run.stderr: raise ValueError('公式转换失败：'+run.stderr[:400])
    with zipfile.ZipFile(out) as z: root=etree.fromstring(z.read('word/document.xml'))
    nodes=root.findall('.//'+qn('m:oMath'))
    if not nodes: raise ValueError('公式没有生成可编辑 OMML：'+latex)
    return [deepcopy(n) for n in nodes]

def write_table(doc,b,number,math_context=None,available_width=None):
    spec=deepcopy(b['table']); rows=spec['rows']; cols=len(rows[0])
    if available_width and sum(spec['widths']) > available_width:
        factor=available_width/sum(spec['widths'])
        spec['widths']=[width*factor for width in spec['widths']]
        spec['width']=available_width
    p=doc.add_paragraph(style='Caption')
    bookmark='table_'+re.sub(r'\W','_',b['id']); p._p.append(element('w:bookmarkStart',id=number,name=bookmark))
    p.add_run('表 ');field(p,' SEQ Table \\* ARABIC ',number);write_runs(p,[], ' '+spec.get('caption',''),math_context);p._p.append(element('w:bookmarkEnd',id=number))
    t=doc.add_table(rows=len(rows),cols=cols);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    width=t._tbl.tblPr.find(qn('w:tblW'))
    if width is None:width=element('w:tblW');t._tbl.tblPr.append(width)
    width.set(qn('w:w'),str(round(spec.get('width',158)*56.693)));width.set(qn('w:type'),'dxa')
    for c,w in enumerate(spec['widths']): t.columns[c].width=Mm(w)
    for ri,row in enumerate(rows):
        t.rows[ri].height=Mm(spec['heights'][ri]); t.rows[ri]._tr.get_or_add_trPr().append(element('w:cantSplit'))
        if ri==0:t.rows[ri]._tr.get_or_add_trPr().append(element('w:tblHeader'))
        for ci,data in enumerate(row):
            if data.get('hidden'):continue
            cell=t.cell(ri,ci)
            rs,cs=data.get('rowspan',1),data.get('colspan',1)
            if rs>1 or cs>1:cell=cell.merge(t.cell(ri+rs-1,ci+cs-1))
            cell.width=Mm(sum(spec['widths'][ci:ci+cs])); cell.text=''
            cell.vertical_alignment={'top':WD_CELL_VERTICAL_ALIGNMENT.TOP,'bottom':WD_CELL_VERTICAL_ALIGNMENT.BOTTOM}.get(data.get('vertical'),WD_CELL_VERTICAL_ALIGNMENT.CENTER)
            p=cell.paragraphs[0];write_runs(p,data.get('runs'),data.get('text',''),math_context);align=data.get('align','auto')
            if align=='auto':align='center' if ri==0 or re.fullmatch(r'[-+]?\d[\d.,% ±Ee-]*',data.get('text','').strip()) else 'left'
            p.alignment={'center':WD_ALIGN_PARAGRAPH.CENTER,'right':WD_ALIGN_PARAGRAPH.RIGHT}.get(align,WD_ALIGN_PARAGRAPH.LEFT)
            if ri==0:
                for r in p.runs:r.bold=True
            borders=element('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                visible=spec.get('style')!='three-line' or edge=='top' and ri==0 or edge=='bottom' and (ri==0 or ri+rs==len(rows))
                borders.append(element('w:'+edge,val='single' if visible else 'nil',sz=round(spec.get('border',1)*8),color=spec.get('color','000000')))
            cell._tc.get_or_add_tcPr().append(borders)
            margins=element('w:tcMar')
            for edge in ['top','left','bottom','right']:margins.append(element('w:'+edge,w=round(spec.get('padding',1.8)*56.693),type='dxa'))
            cell._tc.get_or_add_tcPr().append(margins)

def export(req):
    root=Path(req['directory']);root.mkdir(parents=True,exist_ok=True)
    folder=Path(tempfile.mkdtemp(prefix='manuscript-',dir=root))
    reference=Path(req.get('referenceDoc') or Path(__file__).with_name('reference.docx'))
    if not reference.exists(): raise ValueError('找不到 reference.docx')
    doc=Document(reference)
    if not req.get('referenceDoc'):
        academic_default_styles(doc, EXPORT_HEADERS.get(req.get('operation'), '研究稿'))
    layout=apply_journal_layout(doc, req['project'].get('journalFormat')) if not req.get('referenceDoc') else None
    available_width=((210-2*layout['margin']-(7.5 if layout['columns']>1 else 0))/layout['columns']) if layout else None
    # Retain styles/headers/footers/section properties; discard template example body.
    for node in list(doc.element.body):
        if node.tag!=qn('w:sectPr'):doc.element.body.remove(node)
    p=req['project'];doc.add_heading(p['title'],0);table_number=0
    if layout and layout['columns']>1:
        # A full-width title precedes a continuous two-column article body.
        section_columns(doc.sections[0],1)
        section_columns(doc.add_section(WD_SECTION_START.CONTINUOUS),layout['columns'])
    math_context=dict(pandoc=req.get('pandoc'),folder=folder,warnings=[])
    for section in p['sections']:
        blocks=[b for b in p['blocks'] if b['sectionID']==section['id'] and not (b.get('emptyStarter') and not b.get('text','').strip() and not b.get('materials'))]
        if not blocks:continue
        section_start=len(doc.paragraphs)
        doc.add_heading(section['title'],1)
        for b in blocks:
            kind=b['type'];text=b.get('text','')
            if b.get('pageBreakBefore'):doc.add_page_break()
            if b.get('emptyStarter') and not text.strip():continue
            if kind=='placeholder':raise ValueError('存在待填卡，请补齐后导出')
            if kind=='heading':
                par=doc.add_heading('',2);paragraph_alignment(par,b);write_runs(par,b.get('runs'),text,math_context)
            elif kind=='list':
                for line, parts in list_runs(text,b.get('runs',[])):
                    par=doc.add_paragraph(style='List Number' if b.get('listStyle')=='number' else 'List Bullet')
                    paragraph_alignment(par,b);write_runs(par,parts,line,math_context)
            elif kind=='formula':
                par=doc.add_paragraph()
                for node in formula(text,req.get('pandoc'),folder):par._p.append(node)
            elif kind=='image':
                path=Path(b['imagePath'])
                if not path.is_file():raise ValueError('图片不存在：'+str(path))
                width=b.get('width',140);height=b.get('height')
                if not isinstance(width,(int,float)) or not 0<width<=158:raise ValueError('图片宽度无效')
                if height is not None and (not isinstance(height,(int,float)) or not 0<height<=240):raise ValueError('图片高度无效')
                if available_width and width>available_width:
                    if height is not None:height*=available_width/width
                    width=available_width
                doc.add_picture(str(path),width=Mm(width),**({'height':Mm(height)} if height is not None else {}))
                if text:write_runs(doc.add_paragraph(style='Caption'),b.get('runs'),text,math_context)
            elif kind=='table':table_number+=1;write_table(doc,b,table_number,math_context,available_width)
            elif kind=='citation':
                cited=req.get('citationText',{}).get(str(b.get('citationItemID')))
                if not cited:raise ValueError('引用尚未由 Zotero 格式化')
                doc.add_paragraph(cited)
            elif kind=='crossref':
                target=next((x for x in p['blocks'] if x['id']==b.get('targetID') and x['type']=='table'),None)
                if not target:raise ValueError('交叉引用指向不存在的表格')
                par=doc.add_paragraph();field(par,' REF table_'+re.sub(r'\W','_',target['id'])+' \\h ',text or '表')
            else:
                par=doc.add_paragraph()
                paragraph_alignment(par,b)
                if b.get('spaceBefore'):par.paragraph_format.space_before=Pt(b['spaceBefore'])
                write_runs(par,b.get('runs'),text,math_context)
        if req.get('operation') in KEEP_TASK_OPERATIONS and len(blocks)==4 and re.match(r'^预计[\s：:]',blocks[-1].get('text','')):
            # A plan task's action, deliverable, check and estimated hours form
            # one reading unit. Keep the short task together so the final hours
            # line cannot be orphaned on the next page. Word can still split a
            # task longer than a page; this never inserts artificial page breaks.
            task_paragraphs=doc.paragraphs[section_start:]
            for index,par in enumerate(task_paragraphs):
                par.paragraph_format.keep_together=True
                par.paragraph_format.keep_with_next=index<len(task_paragraphs)-1
    if req.get('bibliography'):
        doc.add_heading('参考文献',1)
        for entry in req['bibliography']:doc.add_paragraph(entry.strip())
    output=folder/'manuscript.docx';doc.save(output)
    with zipfile.ZipFile(output) as z:
        xml=etree.fromstring(z.read('word/document.xml'))
        report=dict(parts=z.namelist(),tables=len(xml.findall('.//'+qn('w:tbl'))),formulas=len(xml.findall('.//'+qn('m:oMath'))),images=len(xml.findall('.//'+qn('w:drawing'))),mathWarnings=math_context['warnings'],textWarnings=math_context.get('textWarnings',[]))
    (folder/'structure-check.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    (folder/'project.json').write_text(json.dumps(p,ensure_ascii=False,indent=2),encoding='utf-8')
    return dict(path=str(output),directory=str(folder),structure=report,citations='Zotero 格式化的静态引用')

if __name__=='__main__':
    try:
        req=json.load(sys.stdin)
        if req['operation']=='manuscript-index':
            from pdf_materials import index_pdf
            value=index_pdf(req)
        elif req['operation']=='manuscript-translation-layout':
            from pdf_materials import translation_layout
            value=translation_layout(req)
        elif req['operation']=='manuscript-plan-export':
            from plan_export import export_plan
            value=export_plan(req)
        elif req['operation']=='manuscript-read-export':
            from read_export import export_read
            value=export_read(req)
        else:
            value=dataset(req['file']) if req['operation']=='manuscript-data' else export(req)
        print(json.dumps(dict(type='result',value=value),ensure_ascii=False))
    except Exception as e:
        print(json.dumps(dict(type='error',message=str(e)),ensure_ascii=False));sys.exit(1)

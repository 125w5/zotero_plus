"""Real DOCX paragraph-format regression; no model or network calls."""
import importlib.util
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

MODULE_PATH = Path(__file__).resolve().parents[2] / 'services/research-engine/manuscript/document.py'
SPEC = importlib.util.spec_from_file_location('manuscript_document', MODULE_PATH)
DOCUMENT = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(DOCUMENT)


class ParagraphAlignmentExport(unittest.TestCase):
    def export(self, blocks, reference_alignment=None):
        with tempfile.TemporaryDirectory() as root:
            reference = Path(root) / 'reference.docx'
            DOCUMENT.default_reference(reference)
            if reference_alignment is not None:
                doc = Document(reference)
                doc.styles['Normal'].paragraph_format.alignment = reference_alignment
                doc.save(reference)
            project = {'title': 'Alignment verification', 'sections': [{'id': 's', 'title': 'Section'}], 'blocks': blocks}
            result = DOCUMENT.export({'directory': root, 'referenceDoc': str(reference), 'project': project})
            return Document(result['path'])

    def test_plain_heading_list_and_rich_text_export_actual_alignment(self):
        blocks = [
            {'id': 'b1', 'sectionID': 's', 'type': 'paragraph', 'text': 'left body', 'textAlign': 'left'},
            {'id': 'b2', 'sectionID': 's', 'type': 'heading', 'text': 'center heading', 'textAlign': 'center'},
            {'id': 'b3', 'sectionID': 's', 'type': 'list', 'text': 'right first\nright second', 'textAlign': 'right'},
            {'id': 'b4', 'sectionID': 's', 'type': 'paragraph', 'text': 'justified xi', 'textAlign': 'justify',
             'runs': [{'text': 'justified ', 'bold': True}, {'text': 'x', 'italic': True}, {'text': 'i', 'sub': True}]}
        ]
        doc = self.export(blocks)
        paragraphs = {paragraph.text: paragraph for paragraph in doc.paragraphs}
        for text, expected in [('left body', WD_ALIGN_PARAGRAPH.LEFT), ('center heading', WD_ALIGN_PARAGRAPH.CENTER),
                               ('right first', WD_ALIGN_PARAGRAPH.RIGHT), ('right second', WD_ALIGN_PARAGRAPH.RIGHT),
                               ('justified xi', WD_ALIGN_PARAGRAPH.JUSTIFY)]:
            self.assertEqual(paragraphs[text].alignment, expected)
        rich = paragraphs['justified xi'].runs
        self.assertTrue(rich[0].bold)
        self.assertTrue(rich[1].italic)
        self.assertTrue(rich[2].font.subscript)
        self.assertEqual(paragraphs['right first'].style.name, 'List Bullet')
        self.assertEqual(paragraphs['right second'].style.name, 'List Bullet')

    def test_older_blocks_keep_their_reference_style(self):
        doc = self.export([{'id': 'old', 'sectionID': 's', 'type': 'paragraph', 'text': 'old paragraph'}], WD_ALIGN_PARAGRAPH.JUSTIFY)
        paragraph = next(paragraph for paragraph in doc.paragraphs if paragraph.text == 'old paragraph')
        self.assertIsNone(paragraph.alignment)
        self.assertEqual(paragraph.style.paragraph_format.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)


if __name__ == '__main__':
    unittest.main()

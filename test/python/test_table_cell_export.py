"""Actual Word table exports use current human text, never obsolete run text."""
from copy import deepcopy
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'services' / 'research-engine' / 'manuscript'))
from document import export, current_text_runs, qn
from lxml import etree
from docx import Document

PANDOC = ROOT.parent / '.tools' / 'pandoc' / 'pandoc-3.11' / 'pandoc.exe'


def project_for(rows):
    return dict(title='单元格正文验证', sections=[dict(id='s', title='结果')], blocks=[
        dict(id='t', sectionID='s', type='table', table=dict(rows=rows, widths=[50] * len(rows[0]), heights=[8] * len(rows), style='grid', caption='人工表格'))
    ])


class TableCellExport(unittest.TestCase):
    def test_old_runs_cannot_export_deleted_text_and_unchanged_prefix_keeps_its_style(self):
        project = project_for([
            [dict(text='人工结果'), dict(text='人工公式')],
            [dict(text='准确率 93%', runs=[dict(text='准确率 91%', bold=True)]),
             dict(text='', runs=[dict(text='不应再出现的旧单元格', italic=True)])]
        ])
        before = deepcopy(project)
        with tempfile.TemporaryDirectory() as folder:
            result = export(dict(directory=folder, project=project))
            doc = Document(result['path'])
            self.assertEqual(doc.tables[0].cell(1, 0).text, '准确率 93%')
            self.assertEqual(doc.tables[0].cell(1, 1).text, '')
            retained = next(run for run in doc.tables[0].cell(1, 0).paragraphs[0].runs if run.text.startswith('准确率'))
            self.assertTrue(retained.bold)
            with zipfile.ZipFile(result['path']) as archive:
                xml = archive.read('word/document.xml').decode('utf-8')
                self.assertNotIn('准确率 91%', xml)
                self.assertNotIn('不应再出现的旧单元格', xml)
            self.assertEqual(len(result['structure']['textWarnings']), 2)
            self.assertTrue(all(row['usedCurrentText'] for row in result['structure']['textWarnings']))
            self.assertEqual(project, before, 'export does not rewrite the human source or its old saved runs')

    def test_consistent_merged_runs_keep_every_character_and_both_original_formats(self):
        runs = [dict(text='结果', bold=True), dict(text=' '), dict(text='变量', italic=True)]
        project = project_for([[dict(text='结果 变量', runs=runs, colspan=2), dict(text='', hidden=True)]])
        with tempfile.TemporaryDirectory() as folder:
            result = export(dict(directory=folder, project=project))
            cell = Document(result['path']).tables[0].cell(0, 0)
            self.assertEqual(cell.text, '结果 变量')
            self.assertTrue(cell.paragraphs[0].runs[0].bold)
            self.assertTrue(cell.paragraphs[0].runs[-1].italic)
            self.assertFalse(result['structure']['textWarnings'])

    @unittest.skipUnless(PANDOC.is_file(), '实际 Pandoc 不可用')
    def test_stale_formula_runs_export_only_the_latest_editable_word_math(self):
        project = project_for([[dict(text='当前变量 $y_j^2$', runs=[dict(text='旧变量 $x_i$', bold=True)])]])
        with tempfile.TemporaryDirectory() as folder:
            result = export(dict(directory=folder, project=project, pandoc=str(PANDOC)))
            with zipfile.ZipFile(result['path']) as archive:
                xml = etree.fromstring(archive.read('word/document.xml'))
            self.assertEqual(len(xml.findall('.//' + qn('m:oMath'))), 1)
            math_text = ''.join(node.text or '' for node in xml.findall('.//' + qn('m:t')))
            self.assertIn('y', math_text)
            self.assertIn('j', math_text)
            self.assertNotIn('x', math_text)
            self.assertNotIn('i', math_text)
            self.assertFalse(xml.findall('.//' + qn('w:drawing')))
            self.assertFalse(result['structure']['mathWarnings'])
            self.assertEqual(len(result['structure']['textWarnings']), 1)

    def test_run_recovery_retains_exact_prefix_and_suffix_without_superseded_characters(self):
        old = [dict(text='测量 ', bold=True), dict(text='91%', italic=True), dict(text=' Hz', sub=True)]
        fixed = current_text_runs(old, '测量 93% Hz')
        self.assertEqual(''.join(row['text'] for row in fixed), '测量 93% Hz')
        self.assertTrue(fixed[0]['bold'])
        self.assertTrue(fixed[-1]['sub'])
        self.assertEqual(current_text_runs(old, ''), [])


if __name__ == '__main__':
    unittest.main()

"""The execution-plan export remains editable and does not overwrite the manuscript."""
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'services' / 'research-engine' / 'manuscript'))
from document import export, qn
from plan_export import export_plan
from lxml import etree
from docx import Document
from docx.shared import RGBColor
from docx.oxml import OxmlElement

PANDOC = ROOT.parent / '.tools' / 'pandoc' / 'pandoc-3.11' / 'pandoc.exe'


def xml_at(path):
    with zipfile.ZipFile(path) as archive:
        return etree.fromstring(archive.read('word/document.xml'))


def math_fixture_project():
    return dict(title='可编辑公式与富文本验证', sections=[dict(id='s', title='方法与核验')], blocks=[
        dict(id='h', sectionID='s', type='heading', text=r'参数 \(\theta_i\) 的定义'),
        dict(id='p', sectionID='s', type='paragraph', text=r'关键变量 \(x_i = \frac{a}{b}\) 与 $y_j$ 是可编辑表达式。', runs=[
            dict(text='关键变量 ', bold=True, color='#314D68'),
            dict(text=r'\(x_', italic=True),
            dict(text=r'i = \frac{a}{b}\)', bold=True),
            dict(text=' 与 ', italic=True, color='#777777'),
            dict(text='$y_j$', bold=True),
            dict(text=' 是可编辑表达式。', fontFamily='Times New Roman', fontSize=11)
        ]),
        dict(id='f', sectionID='s', type='formula', text=r'z=x_i+y_j'),
        dict(id='t', sectionID='s', type='table', table=dict(rows=[
            [dict(text='参数'), dict(text='表达式')],
            [dict(text='权重'), dict(text=r'\(w_i\)')]
        ], widths=[50, 108], heights=[8, 10], caption='参数定义', style='three-line'))
    ])


class ManuscriptPlanExport(unittest.TestCase):
    def test_distinct_editable_documents_preserve_prose_and_plan(self):
        with tempfile.TemporaryDirectory() as directory:
            project = dict(title='人工论文', sections=[dict(id='s', title='方法')], blocks=[dict(id='p', sectionID='s', type='paragraph', text='人工正文保持原样。')])
            manuscript = export(dict(directory=directory, project=project))
            before = Path(manuscript['path']).read_bytes()
            plan = dict(title='人工论文 执行方案', sections=[dict(id='s', title='独立验证')], blocks=[dict(id='p', sectionID='s', type='paragraph', text='按批次划分数据，检查训练与测试没有重叠。')])
            result = export_plan(dict(directory=manuscript['directory'], project=plan))
            self.assertEqual(Path(result['path']).name, 'execution-plan.docx')
            self.assertEqual(Path(result['path']).parent, Path(manuscript['path']).parent)
            self.assertEqual(Path(manuscript['path']).read_bytes(), before)
            with zipfile.ZipFile(result['path']) as archive:
                xml = etree.fromstring(archive.read('word/document.xml'))
                text = ''.join(xml.itertext())
                self.assertIn('按批次划分数据', text)
                self.assertFalse(xml.findall('.//' + qn('w:drawing')))
            with self.assertRaisesRegex(ValueError, '未覆盖'):
                export_plan(dict(directory=manuscript['directory'], project=plan))

    @unittest.skipUnless(PANDOC.is_file(), '实际 Pandoc 不可用')
    def test_actual_pandoc_keeps_math_editable_across_rich_runs_in_both_documents(self):
        with tempfile.TemporaryDirectory() as directory:
            manuscript = export(dict(directory=directory, project=math_fixture_project(), pandoc=str(PANDOC)))
            xml = xml_at(manuscript['path'])
            self.assertEqual(len(xml.findall('.//' + qn('m:oMath'))), 5)
            self.assertEqual(len(xml.findall('.//' + qn('w:drawing'))), 0)
            self.assertFalse(manuscript['structure']['mathWarnings'])
            plain = ''.join(node.text or '' for node in xml.findall('.//' + qn('w:t')))
            self.assertNotIn(r'\frac', plain)
            self.assertNotIn('$y_j$', plain)
            self.assertNotIn(r'\(x_', plain)
            key_run = next(x for x in xml.findall('.//' + qn('w:r')) if x.find(qn('w:t')) is not None and x.find(qn('w:t')).text == '关键变量 ')
            self.assertIsNotNone(key_run.find(qn('w:rPr') + '/' + qn('w:b')))
            self.assertEqual(key_run.find(qn('w:rPr') + '/' + qn('w:color')).get(qn('w:val')), '314D68')
            between = next(x for x in xml.findall('.//' + qn('w:r')) if x.find(qn('w:t')) is not None and x.find(qn('w:t')).text == ' 与 ')
            self.assertIsNotNone(between.find(qn('w:rPr') + '/' + qn('w:i')))
            plan = dict(title='公式核验方案', sections=[dict(id='s', title='实施步骤')], blocks=[dict(id='p', sectionID='s', type='paragraph', text=r'核对 \(x_i=\frac{a}{b}\) 的单位，再验证独立样本。')])
            before = Path(manuscript['path']).read_bytes()
            exported = export_plan(dict(directory=manuscript['directory'], project=plan, pandoc=str(PANDOC)))
            self.assertEqual(len(xml_at(exported['path']).findall('.//' + qn('m:oMath'))), 1)
            self.assertEqual(Path(manuscript['path']).read_bytes(), before)

    def test_failed_inline_conversion_preserves_editable_source_and_reports_warning(self):
        with tempfile.TemporaryDirectory() as directory:
            source = r'人工文字 \(x_i\) 保留'
            project = dict(title='保留来源', sections=[dict(id='s', title='正文')], blocks=[dict(id='p', sectionID='s', type='paragraph', text=source, runs=[dict(text=source, bold=True)])])
            result = export(dict(directory=directory, project=project))
            xml = xml_at(result['path'])
            self.assertIn(source, ''.join(x.text or '' for x in xml.findall('.//' + qn('w:t'))))
            self.assertEqual(len(result['structure']['mathWarnings']), 1)
            self.assertEqual(result['structure']['mathWarnings'][0]['source'], r'\(x_i\)')
            self.assertFalse(xml.findall('.//' + qn('w:drawing')))

    def test_default_academic_styles_and_plan_header_do_not_modify_a_custom_reference(self):
        project = dict(title='样式核验', sections=[dict(id='s', title='方法')], blocks=[dict(id='p', sectionID='s', type='paragraph', text='人工正文。')])
        with tempfile.TemporaryDirectory() as directory:
            manuscript = export(dict(directory=directory, project=project))
            doc = Document(manuscript['path'])
            for name in ['Title', 'Heading 1', 'Heading 2']:
                self.assertEqual(str(doc.styles[name].font.color.rgb), '000000')
                self.assertFalse(doc.styles[name].element.findall('.//' + qn('w:pBdr')))
            self.assertEqual(doc.sections[0].header.paragraphs[0].text, '研究稿')
            plan = export_plan(dict(directory=manuscript['directory'], project=project))
            self.assertEqual(Document(plan['path']).sections[0].header.paragraphs[0].text, '执行方案')
            custom = Document()
            custom.styles['Title'].font.color.rgb = RGBColor.from_string('1D5468')
            custom.styles['Title'].element.get_or_add_pPr().append(OxmlElement('w:pBdr'))
            custom.sections[0].header.paragraphs[0].text = '用户自定页眉'
            reference = Path(directory) / 'custom-reference.docx'
            custom.save(reference)
            retained = export(dict(directory=directory, project=project, referenceDoc=str(reference)))
            doc = Document(retained['path'])
            self.assertEqual(str(doc.styles['Title'].font.color.rgb), '1D5468')
            self.assertTrue(doc.styles['Title'].element.findall('.//' + qn('w:pBdr')))
            self.assertEqual(doc.sections[0].header.paragraphs[0].text, '用户自定页眉')

    @unittest.skipUnless(PANDOC.is_file(), '实际 Pandoc 不可用')
    def test_image_and_numbered_table_captions_keep_their_style_and_editable_math(self):
        import base64
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'figure.png'
            path.write_bytes(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRhsAAAAASUVORK5CYII='))
            project = dict(title='图表图注', sections=[dict(id='s', title='结果')], blocks=[
                dict(id='img', sectionID='s', type='image', imagePath=str(path), text=r'采集结果 \(z_i\)', width=100),
                dict(id='table', sectionID='s', type='table', table=dict(rows=[[dict(text='参数'), dict(text='值')]], widths=[50, 108], heights=[8], caption=r'参数 \(w_i\) 的统计量'))
            ])
            result = export(dict(directory=directory, project=project, pandoc=str(PANDOC)))
            xml = xml_at(result['path'])
            self.assertEqual(len(xml.findall('.//' + qn('m:oMath'))), 2)
            self.assertEqual(len(xml.findall('.//' + qn('w:drawing'))), 1)
            captions = [p for p in xml.findall('.//' + qn('w:p')) if p.find(qn('w:pPr') + '/' + qn('w:pStyle')) is not None and p.find(qn('w:pPr') + '/' + qn('w:pStyle')).get(qn('w:val')) == 'Caption']
            self.assertEqual(len(captions), 2)
            self.assertTrue(all(p.findall('.//' + qn('m:oMath')) for p in captions))
            self.assertIn('SEQ Table', ''.join(x.text or '' for x in captions[1].findall('.//' + qn('w:instrText'))))
            self.assertEqual(len(captions[1].findall('.//' + qn('w:bookmarkStart'))), 1)


if __name__ == '__main__':
    unittest.main()
